from io import BytesIO
import pandas as pd
from openpyxl import load_workbook

def load_excel_file(uploaded_file):
    """
    Har qanday Excel (.xlsx yoki .xls) ni ochib, openpyxl Workbook qaytaradi.
    """
    # Avval .xlsx sifatida o'qishga harakat qilamiz
    try:
        return load_workbook(filename=BytesIO(uploaded_file.read()), data_only=True)
    except Exception:
        # .xls bo'lishi mumkin — pandas bilan o'qiymiz va xlsx'ga aylantiramiz
        uploaded_file.seek(0)
        excel = pd.ExcelFile(uploaded_file)
        buf = BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            for sheet_name in excel.sheet_names:
                df = excel.parse(sheet_name)
                df.to_excel(writer, sheet_name=sheet_name, index=False)
        buf.seek(0)
        return load_workbook(buf, data_only=True)
