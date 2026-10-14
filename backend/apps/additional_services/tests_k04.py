"""K-04 hamkor qidiruvi va saqlangan so'rovlar."""
import io

from django.core.cache import cache
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from apps.accounts.models import User
from apps.additional_services.api_endpoints.PartnersAPI import legal_form
from apps.additional_services.models import IndustrialProduct, IndustrialProductCategory, SavedQuery, Skp
from apps.common.models import Region


class PartnerSearchTests(APITestCase):
    url = "/api/additional-services/partners/search/"

    def setUp(self):
        cache.clear()
        self.tash = Region.objects.create(name="Тошкент ш.", code="1726")
        self.sam = Region.objects.create(name="Самарқанд в.", code="1718")
        root = IndustrialProductCategory.objects.create(name_uz="Ishlab chiqaradigan sanoat")
        self.textile = IndustrialProductCategory.objects.create(name_uz="To'qimachilik", parent=root)
        food = IndustrialProductCategory.objects.create(name_uz="Oziq-ovqat", parent=root)
        skp = Skp.objects.create(code="1")

        def add(inn, name, address, district, cat):
            IndustrialProduct.objects.create(inn=inn, name=name, adress_uz=address, soato_uz=district, category=cat,
                                             skp=skp, phone="+998901234567")

        add("300000001", '"ALFA TEKSTIL" mas`uliyati cheklangan jamiyati', "Toshkent shahri, Chilonzor", "Chilonzor tumani", self.textile)
        add("300000001", '"ALFA TEKSTIL" mas`uliyati cheklangan jamiyati', "Toshkent shahri, Chilonzor", "Chilonzor tumani", food)
        add("300000002", '"BETA" aksiyadorlik jamiyati', "Toshkent shahri, Yunusobod", "Yunusobod tumani", food)
        add("300000003", '"GAMMA TEKSTIL" xususiy korxonasi', "Samarqand viloyati, Urgut", "Urgut tumani", self.textile)
        self.root = root

    def _stirs(self, **params):
        return [r["stir"] for r in self.client.get(self.url, params).data["results"]]

    def test_filters_combine(self):
        self.assertEqual(sorted(self._stirs(region=self.tash.id)), ["300000001", "300000002"])
        self.assertEqual(self._stirs(region=self.tash.id, category=self.textile.id), ["300000001"])
        self.assertEqual(self._stirs(category=self.textile.id, legal_form="XK"), ["300000003"])
        self.assertEqual(self._stirs(q="tekstil", region=self.sam.id), ["300000003"])
        self.assertEqual(self._stirs(district="Yunusobod tumani"), ["300000002"])
        self.assertEqual(sorted(self._stirs(category=self.root.id)), ["300000001", "300000002", "300000003"])  # avlodlar
        self.assertEqual(sorted(self._stirs(q="30000000")), ["300000001", "300000002", "300000003"])  # STIR boshi

    def test_grouped_by_stir_and_no_phone(self):
        data = self.client.get(self.url, {"q": "ALFA"}).data
        self.assertEqual(data["count"], 1)
        row = data["results"][0]
        self.assertEqual(row["products"], 2)
        self.assertEqual(row["region"], "Toshkent shahri")
        self.assertNotIn("phone", row)
        self.assertNotIn("901234567", str(data))

    def test_export_requires_login_and_has_no_personal_columns(self):
        self.assertEqual(self.client.get("/api/additional-services/partners/export/").status_code, 401)
        self.client.force_authenticate(User.objects.create_user(username="exp", password="x"))
        response = self.client.get("/api/additional-services/partners/export/", {"region": self.tash.id})
        self.assertEqual(response.status_code, 200)
        ws = load_workbook(io.BytesIO(response.content)).active
        header = [c.value for c in ws[1]]
        self.assertNotIn("Telefon", " ".join(map(str, header)))
        body = " ".join(str(c.value) for row in ws.iter_rows() for c in row)
        self.assertNotIn("901234567", body)
        self.assertIn("300000002", body)

    def test_map_counts(self):
        data = {r["region_code"]: r["count"] for r in self.client.get("/api/additional-services/partners/map/").data}
        self.assertEqual(data["1726"], 2)
        self.assertEqual(data["1718"], 1)

    def test_legal_form_parsing(self):
        self.assertEqual(legal_form('"X" mas‘uliyati cheklangan jamiyati shaklidagi xorijiy korxonasi')["code"], "XK_FOREIGN")
        self.assertEqual(legal_form('"Y" oilaviy korxonasi')["code"], "OK")
        self.assertEqual(legal_form("Z")["code"], "OTHER")


class SavedQueryTests(APITestCase):
    url = "/api/additional-services/saved-queries/"

    def test_only_own_queries(self):
        a = User.objects.create_user(username="sq_a", password="x")
        b = User.objects.create_user(username="sq_b", password="x")
        SavedQuery.objects.create(user=b, kind="partners", name="B", params={"q": "b"})
        self.client.force_authenticate(a)
        created = self.client.post(self.url, {"kind": "partners", "name": "A", "params": {"region": 5}}, format="json")
        self.assertEqual(created.status_code, 201)
        self.assertEqual([q["name"] for q in self.client.get(self.url, {"kind": "partners"}).data], ["A"])
        foreign = SavedQuery.objects.get(user=b)
        self.assertEqual(self.client.delete(f"{self.url}{foreign.id}/").status_code, 404)
        self.assertEqual(self.client.delete(f"{self.url}{created.data['id']}/").status_code, 204)
